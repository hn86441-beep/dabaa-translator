import os
import io
import json
import base64
import traceback
from typing import Optional
from fastapi import FastAPI, File, UploadFile, Form, HTTPException
from fastapi.responses import StreamingResponse, JSONResponse
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel
import requests
from gtts import gTTS

# استيرادات اختيارية مع التعامل مع الأخطاء
try:
    from deep_translator import GoogleTranslator, MyMemoryTranslator
except:
    GoogleTranslator = MyMemoryTranslator = None

try:
    import pdfplumber
except:
    pdfplumber = None

try:
    import docx
except:
    docx = None

try:
    import openpyxl
except:
    openpyxl = None

try:
    from PIL import Image
except:
    Image = None

# ===== التطبيق =====
app = FastAPI()
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

# ===== المفاتيح (من متغيرات البيئة) =====
GROQ_KEY = os.getenv("GROQ_API_KEY", "").strip()
DEEPL_KEY = os.getenv("DEEPL_API_KEY", "").strip()
COHERE_KEY = os.getenv("COHERE_API_KEY", "").strip()

# ===== اللغات =====
LANGUAGES = {
    "Auto-Detect": {"g": "auto", "d": None, "tts": "en"},
    "Arabic": {"g": "ar", "d": "AR", "tts": "ar"},
    "English": {"g": "en", "d": "EN-US", "tts": "en"},
    "Russian": {"g": "ru", "d": "RU", "tts": "ru"},
    "Chinese": {"g": "zh-CN", "d": "ZH", "tts": "zh-cn"},
    "German": {"g": "de", "d": "DE", "tts": "de"},
    "Spanish": {"g": "es", "d": "ES", "tts": "es"},
    "French": {"g": "fr", "d": "FR", "tts": "fr"},
    "Portuguese": {"g": "pt", "d": "PT-PT", "tts": "pt"},
    "Italian": {"g": "it", "d": "IT", "tts": "it"},
    "Japanese": {"g": "ja", "d": "JA", "tts": "ja"},
    "Korean": {"g": "ko", "d": "KO", "tts": "ko"},
    "Turkish": {"g": "tr", "d": "TR", "tts": "tr"},
    "Dutch": {"g": "nl", "d": "NL", "tts": "nl"},
    "Polish": {"g": "pl", "d": "PL", "tts": "pl"},
    "Ukrainian": {"g": "uk", "d": "UK", "tts": "uk"},
    "Swedish": {"g": "sv", "d": "SV", "tts": "sv"},
    "Danish": {"g": "da", "d": "DA", "tts": "da"},
    "Finnish": {"g": "fi", "d": "FI", "tts": "fi"},
    "Romanian": {"g": "ro", "d": "RO", "tts": "ro"},
    "Hungarian": {"g": "hu", "d": "HU", "tts": "hu"},
    "Czech": {"g": "cs", "d": "CS", "tts": "cs"},
    "Bulgarian": {"g": "bg", "d": "BG", "tts": "bg"},
    "Greek": {"g": "el", "d": "EL", "tts": "el"},
    "Indonesian": {"g": "id", "d": "ID", "tts": "id"},
    "Hindi": {"g": "hi", "d": None, "tts": "hi"},
    "Persian": {"g": "fa", "d": None, "tts": "fa"},
    "Hebrew": {"g": "iw", "d": None, "tts": "iw"},
    "Urdu": {"g": "ur", "d": None, "tts": "ur"},
}

# ===== دوال مساعدة =====
def detect_lang(text):
    if not text: return "en", "English"
    ar = sum(1 for c in text if "\u0600" <= c <= "\u06FF") / max(len(text), 1)
    cy = sum(1 for c in text if "\u0400" <= c <= "\u04FF") / max(len(text), 1)
    cj = sum(1 for c in text if "\u4E00" <= c <= "\u9FFF") / max(len(text), 1)
    if ar > 0.12: return "ar", "Arabic"
    if cy > 0.12: return "ru", "Russian"
    if cj > 0.12: return "zh", "Chinese"
    return "en", "English"

def emotion(text):
    pos = {"شكر","ممتاز","رائع","سعيد","فرح","أحب","جميل","موافق","تمام","حلو",
           "happy","good","great","excellent","love","wonderful","amazing","joy","perfect"}
    neg = {"حزين","سيء","كره","غضب","ألم","خطأ","فشل","مزعج","خطير","قلق",
           "sad","bad","hate","angry","pain","fail","dangerous","terrible","horrible"}
    tl = text.lower()
    p = sum(1 for w in pos if w in tl)
    n = sum(1 for w in neg if w in tl)
    return "😊 إيجابي" if p > n else ("😢 سلبي" if n > p else "😐 محايد")

def quick_domain(text):
    domains = {
        "medical": {"icon": "🏥", "name": "طبي", "keywords": ["doctor","hospital","طبيب","مستشفى","علاج"]},
        "legal": {"icon": "⚖️", "name": "قانوني", "keywords": ["contract","court","law","عقد","قانون"]},
        "political": {"icon": "🏛️", "name": "سياسي", "keywords": ["minister","government","رئيس","وزير","برلمان"]},
        "economic": {"icon": "📈", "name": "اقتصادي", "keywords": ["economic","investment","اقتصاد","استثمار"]},
        "scientific": {"icon": "🔬", "name": "علمي", "keywords": ["research","experiment","بحث","تجربة"]},
        "military": {"icon": "🎖️", "name": "عسكري", "keywords": ["military","army","جيش","عسكري","سلاح"]},
        "sports": {"icon": "⚽", "name": "رياضي", "keywords": ["football","stadium","كرة","ملعب","فريق"]},
        "it": {"icon": "💻", "name": "تقني", "keywords": ["programming","software","برمجة","تطبيق"]},
        "religious": {"icon": "🕌", "name": "ديني", "keywords": ["mosque","prayer","مسجد","صلاة","قرآن"]},
        "literary": {"icon": "📖", "name": "أدبي", "keywords": ["story","poem","قصة","شعر","رواية"]},
    }
    tl = text.lower()
    result = []
    for d, v in domains.items():
        score = sum(tl.count(kw) for kw in v["keywords"])
        if score > 0: result.append((d, score))
    result.sort(key=lambda x: -x[1])
    return [f"{domains[d]['icon']} {domains[d]['name']}" for d, _ in result[:2]]

# ===== الترجمة الذكية =====
def smart_translate(text, target, source="Auto-Detect"):
    if not text or not text.strip():
        return None, "نص فارغ"
    src_g = LANGUAGES.get(source, {}).get("g", "auto")
    tgt_g = LANGUAGES.get(target, {}).get("g", "en")

    # 1) Google (الأسرع والأضمن)
    if GoogleTranslator:
        try:
            res = GoogleTranslator(source=src_g or "auto", target=tgt_g).translate(text)
            if res:
                return res, "Google ⚡"
        except Exception as e:
            pass

    # 2) DeepL
    if DEEPL_KEY:
        info = LANGUAGES.get(target, {})
        if info.get("d"):
            ep = "https://api-free.deepl.com/v2/translate" if DEEPL_KEY.endswith(":fx") else "https://api.deepl.com/v2/translate"
            try:
                r = requests.post(ep, headers={"Authorization": f"DeepL-Auth-Key {DEEPL_KEY}"},
                                  data={"text": text, "target_lang": info["d"]}, timeout=10)
                if r.status_code == 200:
                    return r.json()["translations"][0]["text"], "DeepL ✦"
            except:
                pass

    # 3) Groq AI (للنصوص القصيرة)
    if GROQ_KEY and len(text) <= 1200:
        system = "You are an expert translator. Translate meaning and intent naturally. Return ONLY the translation."
        prompt = f"Translate from {source} to {target}:\n\n{text}"
        try:
            model = "llama-3.1-8b-instant" if len(text) < 300 else "llama-3.3-70b-versatile"
            r = requests.post(
                "https://api.groq.com/openai/v1/chat/completions",
                headers={"Authorization": f"Bearer {GROQ_KEY}", "Content-Type": "application/json"},
                json={"model": model,
                      "messages": [{"role": "system", "content": system}, {"role": "user", "content": prompt}],
                      "max_tokens": 600, "temperature": 0.4},
                timeout=12
            )
            if r.status_code == 200:
                return r.json()["choices"][0]["message"]["content"].strip(), "AI Contextual ✦"
        except:
            pass

    # 4) MyMemory
    if MyMemoryTranslator:
        try:
            s = "en" if (not src_g or src_g == "auto") else src_g
            res = MyMemoryTranslator(source=s, target=tgt_g).translate(text)
            if res:
                return res, "MyMemory"
        except:
            pass
    return None, "فشلت جميع محاولات الترجمة"

# ===== TTS =====
def make_tts(text, lang="en"):
    if not text or not text.strip():
        return None
    try:
        buf = io.BytesIO()
        gTTS(text=text[:500], lang=lang, slow=False).write_to_fp(buf)
        buf.seek(0)
        return buf
    except:
        return None

# ===== استخراج الملفات =====
def extract_file(file_bytes, filename):
    ext = os.path.splitext(filename)[1].lower()
    if ext == ".pdf" and pdfplumber:
        try:
            txt = ""
            with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
                for pg in pdf.pages:
                    t = pg.extract_text()
                    if t: txt += t + "\n"
            return txt.strip() if txt.strip() else None, None
        except Exception as e: return None, str(e)
    if ext == ".docx" and docx:
        try:
            d = docx.Document(io.BytesIO(file_bytes))
            txt = "\n".join(p.text for p in d.paragraphs)
            return txt.strip() if txt.strip() else None, None
        except Exception as e: return None, str(e)
    if ext in (".xlsx", ".xls") and openpyxl:
        try:
            wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
            parts = []
            for sh in wb.worksheets:
                for row in sh.iter_rows():
                    for cell in row:
                        if cell.value is not None:
                            parts.append(str(cell.value))
            txt = "\n".join(parts)
            return txt.strip() if txt.strip() else None, None
        except Exception as e: return None, str(e)
    if ext == ".txt":
        for enc in ("utf-8", "windows-1256", "latin-1"):
            try:
                t = file_bytes.decode(enc)
                if t.strip(): return t.strip(), None
            except: pass
    return None, f"نوع ملف غير مدعوم: {ext}"

# ===== OCR (Groq Vision) =====
def ocr_image(image_bytes):
    if not GROQ_KEY:
        return None, "مفتاح Groq غير موجود"
    try:
        mime = "image/jpeg"
        if Image:
            try:
                fmt = Image.open(io.BytesIO(image_bytes)).format or "JPEG"
                mime = f"image/{'jpeg' if fmt.upper() in ('JPG','JPEG') else fmt.lower()}"
            except: pass
        b64 = base64.b64encode(image_bytes).decode()
        r = requests.post(
            "https://api.groq.com/openai/v1/chat/completions",
            headers={"Authorization": f"Bearer {GROQ_KEY}", "Content-Type": "application/json"},
            json={
                "model": "llama-3.2-11b-vision-preview",
                "temperature": 0,
                "max_tokens": 2048,
                "messages": [{"role": "user", "content": [
                    {"type": "image_url", "image_url": {"url": f"data:{mime};base64,{b64}"}},
                    {"type": "text", "text": "Extract ALL text from this image exactly as written. Preserve original language. Return ONLY the raw text."}
                ]}]
            },
            timeout=25
        )
        if r.status_code == 200:
            txt = r.json()["choices"][0]["message"]["content"].strip()
            return (txt, None) if txt else (None, "لا يوجد نص")
        return None, f"Groq Vision {r.status_code}"
    except Exception as e: return None, str(e)

# ===== STT (Groq Whisper) =====
def groq_stt(audio_bytes, lang="auto"):
    if not GROQ_KEY:
        return None, "مفتاح Groq غير موجود"
    lc = None
    if lang and lang != "auto":
        lc = {"zh-CN": "zh", "zh-cn": "zh", "iw": "he"}.get(lang, lang[:2])
    files = {"file": ("audio.wav", audio_bytes, "audio/wav"), "model": (None, "whisper-large-v3-turbo")}
    if lc: files["language"] = (None, lc)
    try:
        r = requests.post("https://api.groq.com/openai/v1/audio/transcriptions",
                          headers={"Authorization": f"Bearer {GROQ_KEY}"}, files=files, timeout=30)
        if r.status_code == 200:
            return r.json(), None
        return None, f"Groq STT {r.status_code}"
    except Exception as e: return None, str(e)

# ===== الترجمات (Subtitles) =====
def seconds_to_srt(s):
    h = int(s // 3600); m = int((s % 3600) // 60); sc = int(s % 60); ms = int((s % 1) * 1000)
    return f"{h:02d}:{m:02d}:{sc:02d},{ms:03d}"

def generate_subtitles(audio_bytes, target_lang):
    data, err = groq_stt(audio_bytes, "auto")
    if err: return None, err
    segments = data.get("segments", [])
    if not segments: return None, "لم يُكتشف كلام"
    trans_blocks = []
    for i, seg in enumerate(segments, 1):
        start = seconds_to_srt(seg.get("start", 0))
        end = seconds_to_srt(seg.get("end", 0))
        text = seg.get("text", "").strip()
        translated, _ = smart_translate(text, target_lang, "Auto-Detect")
        trans_blocks.append({"index": i, "start": start, "end": end, "text": translated or text})
    def to_srt(blocks):
        return "\n\n".join(f"{b['index']}\n{b['start']} --> {b['end']}\n{b['text']}" for b in blocks)
    return {"segments": trans_blocks, "srt_translated": to_srt(trans_blocks)}, None

# ===== مجموعة (Group) =====
def group_chat(audio_bytes, target_lang):
    data, err = groq_stt(audio_bytes, "auto")
    if err: return None, err
    segments = data.get("segments", [])
    if not segments: return None, "لم يُكتشف كلام"
    groups, cur = [], []
    for seg in segments:
        if not cur: cur.append(seg)
        else:
            if seg.get("start", 0) - cur[-1].get("end", 0) >= 0.6:
                groups.append(cur); cur = []
            cur.append(seg)
    if cur: groups.append(cur)
    results, spk_map, cnt = [], {}, 1
    icons = ["🧑", "👤", "👩", "👱", "🧔", "🧕", "👲", "🧑‍💼", "👩‍💼", "🙍"]
    for grp in groups:
        text = " ".join(s.get("text", "").strip() for s in grp).strip()
        if not text: continue
        lc, ln = detect_lang(text)
        if lc not in spk_map: spk_map[lc] = cnt; cnt += 1
        spk = spk_map[lc]
        translated, engine = smart_translate(text, target_lang, ln)
        results.append({
            "speaker": spk,
            "icon": icons[(spk - 1) % len(icons)],
            "language": ln,
            "original": text,
            "translated": translated or text,
            "engine": engine,
            "timestamp": f"{grp[0].get('start',0):.1f}s – {grp[-1].get('end',0):.1f}s",
            "color": f"hsl({(spk*83+140)%360}, 65%, 50%)"
        })
    return {"segments": results}, None

# ===== نماذج الطلبات =====
class TranslateRequest(BaseModel):
    text: str
    source: str = "Auto-Detect"
    target: str = "Arabic"

class TTSRequest(BaseModel):
    text: str
    lang: str = "en"

# ===== نقاط النهاية (API) =====
@app.get("/api/health")
async def health():
    return {"status": "ok", "groq": bool(GROQ_KEY), "deepl": bool(DEEPL_KEY)}

@app.post("/api/translate")
async def translate_endpoint(req: TranslateRequest):
    try:
        if not req.text or not req.text.strip():
            raise HTTPException(status_code=400, detail="نص فارغ")
        translation, engine = smart_translate(req.text, req.target, req.source)
        if translation is None:
            raise HTTPException(status_code=500, detail="فشلت الترجمة")
        return {
            "translation": translation,
            "emotion": emotion(req.text),
            "engine": engine,
            "domains": quick_domain(req.text)
        }
    except Exception as e:
        traceback.print_exc()
        raise HTTPException(status_code=500, detail=str(e))

@app.post("/api/tts")
async def tts_endpoint(req: TTSRequest):
    try:
        audio = make_tts(req.text, req.lang)
        if audio is None:
            raise HTTPException(status_code=500, detail="فشل توليد الصوت")
        return StreamingResponse(audio, media_type="audio/mpeg")
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))

@app.post("/api/extract")
async def extract_endpoint(file: UploadFile = File(...)):
    try:
        content = await file.read()
        text, err = extract_file(content, file.filename)
        if err: raise HTTPException(status_code=400, detail=err)
        if not text: raise HTTPException(status_code=400, detail="لا يوجد نص في الملف")
        return {"text": text}
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))

@app.post("/api/ocr")
async def ocr_endpoint(image: UploadFile = File(...)):
    try:
        img = await image.read()
        text, err = ocr_image(img)
        if err: raise HTTPException(status_code=500, detail=err)
        return {"text": text}
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))

@app.post("/api/subtitle")
async def subtitle_endpoint(file: UploadFile = File(...), target: str = Form("Arabic")):
    try:
        audio = await file.read()
        result, err = generate_subtitles(audio, target)
        if err: raise HTTPException(status_code=500, detail=err)
        return result
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))

@app.post("/api/group")
async def group_endpoint(file: UploadFile = File(...), target: str = Form("Arabic")):
    try:
        audio = await file.read()
        result, err = group_chat(audio, target)
        if err: raise HTTPException(status_code=500, detail=err)
        return result
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))

@app.get("/")
async def root():
    return {"message": "HN Translator API is running"}
