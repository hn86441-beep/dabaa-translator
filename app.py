import streamlit as st
st.set_page_config(page_title="HN TRANSLATOR", page_icon="🌐", layout="centered")

import requests, os, json, tempfile, io, base64, sqlite3, time
from requests_toolbelt.multipart.encoder import MultipartEncoder
from collections import OrderedDict
from datetime import datetime
from PIL import Image

try:
    import pdfplumber;     PDF_OK = True
except ImportError:        PDF_OK = False
try:
    import docx;           DOCX_OK = True
except ImportError:        DOCX_OK = False
try:
    import openpyxl;       EXCEL_OK = True
except ImportError:        EXCEL_OK = False
try:
    from transformers import pipeline as hf_pipeline; HF_OK = True
except ImportError:        HF_OK = False
try:
    from langdetect import detect as _ld, DetectorFactory
    DetectorFactory.seed = 42; LANGDETECT_OK = True
except ImportError:        LANGDETECT_OK = False
try:
    import easyocr;        EASYOCR_OK = True
except ImportError:        EASYOCR_OK = False

from deep_translator import GoogleTranslator, MyMemoryTranslator
from gtts import gTTS

# ═══════════════════════════════════════════════════════════
#  جدول اللغات
# ═══════════════════════════════════════════════════════════
ALL_LANGUAGES = {
    "Auto-Detect": {"google":"auto",  "deepl":None,    "gtts":"en",    "whisper":None,  "ocr":["en","ar"]},
    "Arabic":      {"google":"ar",    "deepl":"AR",    "gtts":"ar",    "whisper":"ar",  "ocr":["ar","en"]},
    "English":     {"google":"en",    "deepl":"EN-US", "gtts":"en",    "whisper":"en",  "ocr":["en"]},
    "Russian":     {"google":"ru",    "deepl":"RU",    "gtts":"ru",    "whisper":"ru",  "ocr":["ru","en"]},
    "Chinese":     {"google":"zh-CN", "deepl":"ZH",    "gtts":"zh-cn", "whisper":"zh",  "ocr":["ch_sim","en"]},
    "German":      {"google":"de",    "deepl":"DE",    "gtts":"de",    "whisper":"de",  "ocr":["de","en"]},
    "Spanish":     {"google":"es",    "deepl":"ES",    "gtts":"es",    "whisper":"es",  "ocr":["es","en"]},
    "French":      {"google":"fr",    "deepl":"FR",    "gtts":"fr",    "whisper":"fr",  "ocr":["fr","en"]},
    "Portuguese":  {"google":"pt",    "deepl":"PT-PT", "gtts":"pt",    "whisper":"pt",  "ocr":["pt","en"]},
    "Italian":     {"google":"it",    "deepl":"IT",    "gtts":"it",    "whisper":"it",  "ocr":["it","en"]},
    "Japanese":    {"google":"ja",    "deepl":"JA",    "gtts":"ja",    "whisper":"ja",  "ocr":["ja","en"]},
    "Korean":      {"google":"ko",    "deepl":"KO",    "gtts":"ko",    "whisper":"ko",  "ocr":["ko","en"]},
    "Turkish":     {"google":"tr",    "deepl":"TR",    "gtts":"tr",    "whisper":"tr",  "ocr":["tr","en"]},
    "Dutch":       {"google":"nl",    "deepl":"NL",    "gtts":"nl",    "whisper":"nl",  "ocr":["nl","en"]},
    "Polish":      {"google":"pl",    "deepl":"PL",    "gtts":"pl",    "whisper":"pl",  "ocr":["pl","en"]},
    "Ukrainian":   {"google":"uk",    "deepl":"UK",    "gtts":"uk",    "whisper":"uk",  "ocr":["uk","en"]},
    "Swedish":     {"google":"sv",    "deepl":"SV",    "gtts":"sv",    "whisper":"sv",  "ocr":["sv","en"]},
    "Danish":      {"google":"da",    "deepl":"DA",    "gtts":"da",    "whisper":"da",  "ocr":["da","en"]},
    "Finnish":     {"google":"fi",    "deepl":"FI",    "gtts":"fi",    "whisper":"fi",  "ocr":["fi","en"]},
    "Romanian":    {"google":"ro",    "deepl":"RO",    "gtts":"ro",    "whisper":"ro",  "ocr":["ro","en"]},
    "Hungarian":   {"google":"hu",    "deepl":"HU",    "gtts":"hu",    "whisper":"hu",  "ocr":["hu","en"]},
    "Czech":       {"google":"cs",    "deepl":"CS",    "gtts":"cs",    "whisper":"cs",  "ocr":["cs","en"]},
    "Bulgarian":   {"google":"bg",    "deepl":"BG",    "gtts":"bg",    "whisper":"bg",  "ocr":["bg","en"]},
    "Greek":       {"google":"el",    "deepl":"EL",    "gtts":"el",    "whisper":"el",  "ocr":["el","en"]},
    "Indonesian":  {"google":"id",    "deepl":"ID",    "gtts":"id",    "whisper":"id",  "ocr":["id","en"]},
    "Hindi":       {"google":"hi",    "deepl":None,    "gtts":"hi",    "whisper":"hi",  "ocr":["hi","en"]},
    "Persian":     {"google":"fa",    "deepl":None,    "gtts":"fa",    "whisper":"fa",  "ocr":["fa","en"]},
    "Hebrew":      {"google":"iw",    "deepl":None,    "gtts":"iw",    "whisper":"he",  "ocr":["he","en"]},
    "Urdu":        {"google":"ur",    "deepl":None,    "gtts":"ur",    "whisper":"ur",  "ocr":["ur","en"]},
}
LANG_NAMES         = list(ALL_LANGUAGES.keys())
LANG_NAMES_NO_AUTO = [k for k in LANG_NAMES if k != "Auto-Detect"]

_LC2NAME = {
    "ar":"Arabic","en":"English","ru":"Russian","zh-cn":"Chinese","zh-tw":"Chinese","zh":"Chinese",
    "de":"German","es":"Spanish","fr":"French","pt":"Portuguese","it":"Italian",
    "ja":"Japanese","ko":"Korean","tr":"Turkish","nl":"Dutch","pl":"Polish",
    "uk":"Ukrainian","sv":"Swedish","da":"Danish","fi":"Finnish","ro":"Romanian",
    "hu":"Hungarian","cs":"Czech","bg":"Bulgarian","el":"Greek","id":"Indonesian",
    "hi":"Hindi","fa":"Persian","he":"Hebrew","ur":"Urdu",
}

# ═══════════════════════════════════════════════════════════
#  قاعدة البيانات
# ═══════════════════════════════════════════════════════════
_DB = os.path.join(tempfile.gettempdir(), "hn_translations.db")

def _db():
    return sqlite3.connect(_DB, timeout=10, check_same_thread=False)

def init_db():
    try:
        c = _db()
        c.execute('''CREATE TABLE IF NOT EXISTS history (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            original TEXT NOT NULL DEFAULT '',
            translated TEXT NOT NULL DEFAULT '',
            emotion TEXT DEFAULT '',
            source_lang TEXT DEFAULT '',
            target_lang TEXT DEFAULT '',
            engine TEXT DEFAULT '',
            timestamp TEXT DEFAULT '')''')
        ex = {r[1] for r in c.execute("PRAGMA table_info(history)")}
        for col, typ in [("engine","TEXT DEFAULT ''"),("emotion","TEXT DEFAULT ''"),
                         ("source_lang","TEXT DEFAULT ''"),("target_lang","TEXT DEFAULT ''")]:
            if col not in ex:
                c.execute(f"ALTER TABLE history ADD COLUMN {col} {typ}")
        c.commit(); c.close()
    except Exception: pass

def _mem(entry):
    if "_mem" not in st.session_state: st.session_state["_mem"] = []
    st.session_state["_mem"].insert(0, entry)
    st.session_state["_mem"] = st.session_state["_mem"][:300]

def save_translation(orig, trans, emotion, src, tgt, engine=""):
    e = {"original":str(orig or "")[:500],"translated":str(trans or "")[:500],
         "emotion":str(emotion or ""),"source_lang":str(src or ""),
         "target_lang":str(tgt or ""),"engine":str(engine or ""),
         "time":datetime.now().strftime("%Y-%m-%d %H:%M")}
    _mem(e)
    try:
        c = _db()
        c.execute('''INSERT INTO history (original,translated,emotion,source_lang,target_lang,engine,timestamp)
                     VALUES (?,?,?,?,?,?,?)''',
                  tuple(e[k] for k in ["original","translated","emotion","source_lang","target_lang","engine","time"]))
        c.commit(); c.close()
    except Exception: pass

def get_history(limit=100):
    try:
        c = _db()
        rows = c.execute('''SELECT original,translated,emotion,source_lang,target_lang,engine,timestamp
                            FROM history ORDER BY id DESC LIMIT ?''',(limit,)).fetchall()
        c.close()
        if rows:
            return [{"original":r[0],"translated":r[1],"emotion":r[2] or "",
                     "source_lang":r[3] or "","target_lang":r[4] or "",
                     "engine":r[5] or "","time":r[6] or ""} for r in rows]
    except Exception: pass
    return st.session_state.get("_mem",[])[:limit]

def clear_history():
    st.session_state["_mem"] = []
    try:
        c = _db(); c.execute("DELETE FROM history"); c.commit(); c.close()
    except Exception: pass

init_db()

# ═══════════════════════════════════════════════════════════
#  تحليل المشاعر
# ═══════════════════════════════════════════════════════════
@st.cache_resource(show_spinner=False)
def _load_emotion():
    if not HF_OK: return None
    try:
        return hf_pipeline("text-classification",
                           model="tabularisai/multilingual-sentiment-analysis",
                           top_k=None, truncation=True, max_length=512)
    except Exception: return None

_EM = _load_emotion()

_POS = ["شكر","ممتاز","رائع","سعيد","فرح","أحب","جميل","موافق","صحيح",
        "happy","good","great","excellent","love","wonderful","amazing","joy","perfect","awesome"]
_NEG = ["حزين","سيء","كره","غضب","ألم","خطأ","فشل","مزعج","لا","خطير","قلق","توتر",
        "sad","bad","hate","angry","pain","error","fail","no","dangerous","terrible","horrible","awful"]

def analyze_emotion(text: str) -> str:
    if not text or not text.strip(): return "😐 محايد"
    if _EM is not None:
        try:
            preds = _EM(text[:512])
            items = preds[0] if isinstance(preds[0], list) else preds
            best  = max(items, key=lambda x: x.get("score", 0))
            lbl, sc = best.get("label","").upper(), best.get("score",0)
            if sc >= 0.50:
                if any(x in lbl for x in ("POSITIVE","POS","HAPPY","JOY","VERY_POSITIVE")): return "😊 إيجابي"
                if any(x in lbl for x in ("NEGATIVE","NEG","SAD","ANGER","VERY_NEGATIVE")): return "😔 سلبي"
                if "NEUTRAL" in lbl: return "😐 محايد"
        except Exception: pass
    tl = text.lower()
    pos = sum(1 for w in _POS if w in tl)
    neg = sum(1 for w in _NEG if w in tl)
    if pos > neg: return "😊 إيجابي"
    if neg > pos: return "😔 سلبي"
    return "😐 محايد"

# ═══════════════════════════════════════════════════════════
#  كشف لغة النص
# ═══════════════════════════════════════════════════════════
def detect_lang(text: str) -> tuple:
    if not text: return "en","English"
    ar  = sum(1 for c in text if "\u0600"<=c<="\u06FF") / max(len(text),1)
    cyr = sum(1 for c in text if "\u0400"<=c<="\u04FF") / max(len(text),1)
    cjk = sum(1 for c in text if "\u4E00"<=c<="\u9FFF") / max(len(text),1)
    if ar  > 0.15: return "ar","Arabic"
    if cyr > 0.15: return "ru","Russian"
    if cjk > 0.15: return "zh","Chinese"
    if LANGDETECT_OK and len(text.strip()) >= 5:
        try:
            code = _ld(text)
            return code, _LC2NAME.get(code, code.capitalize())
        except Exception: pass
    return "en","English"

# ═══════════════════════════════════════════════════════════
#  الترجمة
# ═══════════════════════════════════════════════════════════
def _deepl(text, tgt_code, ak):
    ep = ("https://api-free.deepl.com/v2/translate" if ak.endswith(":fx")
          else "https://api.deepl.com/v2/translate")
    try:
        r = requests.post(ep, headers={"Authorization":f"DeepL-Auth-Key {ak}"},
                          data={"text":text,"target_lang":tgt_code}, timeout=15)
        return (r.json()["translations"][0]["text"],None) if r.status_code==200 else (None,f"DeepL {r.status_code}")
    except Exception as e: return None,str(e)

def _google(text, tgt, src="auto"):
    try: return GoogleTranslator(source=src or "auto",target=tgt).translate(text),None
    except Exception as e1:
        try:
            s = "en" if (not src or src=="auto") else src
            return MyMemoryTranslator(source=s,target=tgt).translate(text),None
        except Exception as e2: return None,f"{e1}|{e2}"

def translate_text(text:str, tgt_name:str, src_name:str="Auto-Detect") -> tuple:
    if not text or not text.strip(): return None,"no text"
    info = ALL_LANGUAGES.get(tgt_name,{})
    ak = st.session_state.get("deepl_api_key","")
    if ak and info.get("deepl"):
        r,_ = _deepl(text, info["deepl"], ak)
        if r: return r,"DeepL ✦"
    src_g = ALL_LANGUAGES.get(src_name,{}).get("google","auto")
    r,err = _google(text, info.get("google","en"), src_g)
    return (r,"Google") if r else (None, err or "فشلت الترجمة")

# ═══════════════════════════════════════════════════════════
#  TTS
# ═══════════════════════════════════════════════════════════
def tts(text:str, lang:str="en") -> io.BytesIO | None:
    if not text or not text.strip(): return None
    try:
        buf = io.BytesIO()
        gTTS(text=text, lang=lang, slow=False).write_to_fp(buf)
        buf.seek(0); return buf
    except Exception: return None

# ═══════════════════════════════════════════════════════════
#  استخراج النص من الملفات
# ═══════════════════════════════════════════════════════════
def extract_file(file_bytes:bytes, filename:str) -> tuple:
    ext = os.path.splitext(filename)[1].lower()
    if ext==".pdf":
        if not PDF_OK: return None,"pdfplumber غير مثبت"
        try:
            txt=""
            with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
                for pg in pdf.pages:
                    t=pg.extract_text()
                    if t: txt+=t+"\n"
            return (txt.strip(),None) if txt.strip() else (None,"لا نص في PDF")
        except Exception as e: return None,str(e)
    if ext==".docx":
        if not DOCX_OK: return None,"python-docx غير مثبت"
        try:
            d=docx.Document(io.BytesIO(file_bytes))
            txt="\n".join(p.text for p in d.paragraphs)
            return (txt.strip(),None) if txt.strip() else (None,"لا نص في DOCX")
        except Exception as e: return None,str(e)
    if ext in (".xlsx",".xls"):
        if not EXCEL_OK: return None,"openpyxl غير مثبت"
        try:
            wb=openpyxl.load_workbook(io.BytesIO(file_bytes),data_only=True)
            parts=[str(c.value) for sh in wb.worksheets for row in sh.iter_rows()
                   for c in row if c.value is not None]
            txt="\n".join(parts)
            return (txt.strip(),None) if txt.strip() else (None,"لا نص في Excel")
        except Exception as e: return None,str(e)
    if ext==".txt":
        for enc in ("utf-8","windows-1256","latin-1"):
            try:
                txt=file_bytes.decode(enc)
                if txt.strip(): return txt.strip(),None
            except Exception: pass
        return None,"تعذر قراءة الملف"
    return None,f"نوع غير مدعوم: {ext}"

# ═══════════════════════════════════════════════════════════
#  OCR — Groq Vision أولاً، EasyOCR احتياطي
# ═══════════════════════════════════════════════════════════
def _mime(img_bytes:bytes) -> str:
    try:
        fmt=Image.open(io.BytesIO(img_bytes)).format or "JPEG"
        return f"image/{'jpeg' if fmt.upper() in ('JPG','JPEG') else fmt.lower()}"
    except Exception: return "image/jpeg"

def ocr_groq(img_bytes:bytes) -> tuple:
    ak=st.session_state.get("groq_api_key","")
    if not ak: return None,"مفتاح Groq غير موجود"
    try:
        mime=_mime(img_bytes); b64=base64.b64encode(img_bytes).decode()
        for model in ["meta-llama/llama-4-scout-17b-16e-instruct","llama-3.2-11b-vision-preview"]:
            r=requests.post("https://api.groq.com/openai/v1/chat/completions",
                headers={"Authorization":f"Bearer {ak}","Content-Type":"application/json"},
                json={"model":model,"temperature":0,"max_tokens":2048,
                      "messages":[{"role":"user","content":[
                          {"type":"image_url","image_url":{"url":f"data:{mime};base64,{b64}"}},
                          {"type":"text","text":
                           "Extract ALL text from this image exactly as written. "
                           "Keep the original language (Arabic stays Arabic, etc). "
                           "Return ONLY the raw text — no labels, no explanations."}
                      ]}]}, timeout=30)
            if r.status_code==200:
                txt=r.json()["choices"][0]["message"]["content"].strip()
                return (txt,None) if txt else (None,"لم يُعثر على نص")
            if r.status_code==404: continue
            return None,f"Groq Vision {r.status_code}"
        return None,"النماذج غير متاحة"
    except Exception as e: return None,str(e)

@st.cache_resource(show_spinner=False)
def _easyocr(langs_t:tuple):
    if not EASYOCR_OK: return None
    try: return easyocr.Reader(list(langs_t),gpu=False,verbose=False)
    except Exception: return None

def ocr_image(img_bytes:bytes, ocr_langs:list) -> tuple:
    if st.session_state.get("groq_api_key"):
        txt,err=ocr_groq(img_bytes)
        if txt: return txt,None
    if not EASYOCR_OK: return None,"EasyOCR غير مثبت"
    try:
        import numpy as np
        reader=_easyocr(tuple(ocr_langs))
        if not reader: return None,"تعذر تحميل EasyOCR"
        arr=np.array(Image.open(io.BytesIO(img_bytes)).convert("RGB"))
        result=reader.readtext(arr,detail=0,paragraph=True)
        txt=" ".join(result).strip()
        return (txt,None) if txt else (None,"لم يُعثر على نص")
    except Exception as e: return None,str(e)

# ═══════════════════════════════════════════════════════════
#  STT — Groq → Cohere → Whisper محلي
# ═══════════════════════════════════════════════════════════
def _wl(code:str):
    if not code or code=="auto": return None
    return {"zh-CN":"zh","zh-cn":"zh","iw":"he"}.get(code, code[:2])

def _groq_stt(audio:bytes, lang="auto", verbose=False) -> tuple:
    ak=st.session_state.get("groq_api_key","")
    if not ak: return None,"مفتاح Groq غير موجود"
    lc=_wl(lang)
    files={"file":("audio.wav",audio,"audio/wav"),
           "model":(None,"whisper-large-v3-turbo"),
           "response_format":(None,"verbose_json" if verbose else "json")}
    if verbose: files["timestamp_granularities[]"]=(None,"segment")
    if lc: files["language"]=(None,lc)
    try:
        r=requests.post("https://api.groq.com/openai/v1/audio/transcriptions",
                        headers={"Authorization":f"Bearer {ak}"},files=files,timeout=60)
        if r.status_code==200:
            data=r.json()
            if verbose: return data,None
            txt=data.get("text","").strip()
            return (txt,None) if txt else (None,"لم يُكتشف كلام")
        return None,f"Groq {r.status_code}"
    except Exception as e: return None,str(e)

def _cohere_stt(audio:bytes, lang="en") -> tuple:
    ak=st.session_state.get("cohere_api_key","")
    if not ak: return None,"مفتاح Cohere غير موجود"
    lc=_wl(lang) or "en"
    try:
        fields=OrderedDict([("language",lc),("model","cohere-transcribe-03-2026"),
                             ("file",("audio.wav",audio,"audio/wav"))])
        enc=MultipartEncoder(fields=fields)
        r=requests.post("https://api.cohere.com/v2/audio/transcriptions",
                        headers={"Authorization":f"Bearer {ak}","Content-Type":enc.content_type},
                        data=enc,timeout=30)
        if r.status_code==200:
            txt=r.json().get("text","").strip()
            return (txt,None) if txt else (None,"لم يُكتشف كلام")
        return None,f"Cohere {r.status_code}"
    except Exception as e: return None,str(e)

@st.cache_resource(show_spinner=False)
def _load_whisper():
    try:
        from faster_whisper import WhisperModel
        return WhisperModel("small",device="cpu",compute_type="int8")
    except Exception: return None

def _local_stt(audio:bytes, lang=None) -> tuple:
    m=_load_whisper()
    if not m: return None,"Whisper غير متاح"
    tmp=None
    try:
        with tempfile.NamedTemporaryFile(delete=False,suffix=".wav") as f:
            f.write(audio); tmp=f.name
        segs,_=m.transcribe(tmp,language=lang,beam_size=5,vad_filter=True)
        txt=" ".join(s.text for s in segs).strip()
        return (txt,None) if txt else (None,"لم يُكتشف كلام")
    except Exception as e: return None,str(e)
    finally:
        if tmp and os.path.exists(tmp): os.unlink(tmp)

def stt(audio:bytes, lang_code:str="auto") -> tuple:
    wl=_wl(lang_code)
    if st.session_state.get("groq_api_key"):
        r,_=_groq_stt(audio,lang_code)
        if r: return r,None
    if st.session_state.get("cohere_api_key"):
        r,_=_cohere_stt(audio,lang_code)
        if r: return r,None
    return _local_stt(audio,lang=wl)

# ═══════════════════════════════════════════════════════════
#  محادثة جماعية
# ═══════════════════════════════════════════════════════════
_ICONS=["🧑","👤","👩","👱","🧔","🧕","👲","🧑‍💼","👩‍💼","🙍"]

def _grp_segs(segs:list,gap=0.6) -> list:
    if not segs: return []
    groups,cur=[],[segs[0]]
    for i in range(1,len(segs)):
        if segs[i].get("start",0)-segs[i-1].get("end",0)>=gap:
            groups.append(cur); cur=[]
        cur.append(segs[i])
    if cur: groups.append(cur)
    return groups

def group_analyze(audio:bytes, tgt:str) -> tuple:
    tgt_gtts=ALL_LANGUAGES.get(tgt,{}).get("gtts","en")

    def _make(txt,lc,ln,t0=0,t1=0):
        sn=_LC2NAME.get(lc,"Auto-Detect")
        tr,eng=translate_text(txt,tgt,sn)
        return [{"speaker_num":1,"icon":_ICONS[0],"lc":lc,"lang":ln,
                 "orig":txt,"trans":tr or "","engine":eng,"gtts":tgt_gtts,"t0":t0,"t1":t1}]

    data,err=_groq_stt(audio,lang="auto",verbose=True)
    if data is None:
        txt=None
        if st.session_state.get("cohere_api_key"): txt,_=_cohere_stt(audio)
        if not txt: txt,err=_local_stt(audio)
        if not txt: return None,err or "تعذر التعرف"
        lc,ln=detect_lang(txt)
        return _make(txt,lc,ln), None

    segs=data.get("segments",[])
    full=data.get("text","").strip()
    if not segs:
        if not full: return None,"لم يُكتشف كلام"
        lc,ln=detect_lang(full)
        return _make(full,lc,ln), None

    groups=_grp_segs(segs,0.6)
    results=[]; lang2spk={}; prev_lc=None; spk_n=1

    for grp in groups:
        txt=" ".join(s.get("text","").strip() for s in grp).strip()
        if not txt: continue
        lc,ln=detect_lang(txt)
        if lc not in lang2spk:
            lang2spk[lc]=spk_n; spk_n+=1
        spk=lang2spk[lc]
        icon=_ICONS[min(spk-1,len(_ICONS)-1)]
        sn=_LC2NAME.get(lc,"Auto-Detect")
        tr,eng=translate_text(txt,tgt,sn)
        results.append({"speaker_num":spk,"icon":icon,"lc":lc,"lang":ln,
                         "orig":txt,"trans":tr or "","engine":eng,"gtts":tgt_gtts,
                         "t0":grp[0].get("start",0),"t1":grp[-1].get("end",0)})
        prev_lc=lc

    return (results,None) if results else (None,"لم يُكتشف كلام")

# ═══════════════════════════════════════════════════════════
#  كشف مجال النص
# ═══════════════════════════════════════════════════════════
_DOM={
    "political": {"e":"🏛️","n":"Political","kw":["minister","government","president","وزير","حكومة","رئيس"]},
    "legal":     {"e":"⚖️","n":"Legal",    "kw":["contract","court","law","عقد","قانون","محكمة"]},
    "economic":  {"e":"📈","n":"Economic", "kw":["economic","investment","اقتصاد","استثمار"]},
    "medical":   {"e":"🏥","n":"Medical",  "kw":["doctor","hospital","طبيب","مستشفى","علاج"]},
    "scientific":{"e":"🔬","n":"Scientific","kw":["research","experiment","بحث","تجربة"]},
    "military":  {"e":"🎖️","n":"Military","kw":["military","army","جيش","عسكري","سلاح"]},
    "sports":    {"e":"⚽","n":"Sports",   "kw":["football","stadium","كرة","ملعب","فريق"]},
    "it":        {"e":"💻","n":"IT/Tech",  "kw":["programming","software","برمجة","موقع","تطبيق"]},
    "tourism":   {"e":"✈️","n":"Tourism",  "kw":["hotel","travel","فندق","سفر","مطار"]},
}
def detect_dom(text:str) -> list:
    tl=text.lower()
    return sorted((d for d,v in _DOM.items() if any(k in tl for k in v["kw"])),
                  key=lambda d:-sum(tl.count(k) for k in _DOM[d]["kw"]))[:3]

# ═══════════════════════════════════════════════════════════
#  Session State
# ═══════════════════════════════════════════════════════════
def _sec(k,d=""):
    try: return st.secrets.get(k,d) or d
    except Exception: return d

for k,v in {"theme":"dark","groq_api_key":_sec("GROQ_API_KEY"),
             "deepl_api_key":_sec("DEEPL_API_KEY"),"cohere_api_key":_sec("COHERE_API_KEY"),
             "src_lang":"Auto-Detect","tgt_lang":"Arabic","input_text":""}.items():
    if k not in st.session_state: st.session_state[k]=v

# ═══════════════════════════════════════════════════════════
#  CSS — خط واضح + تصميم محسّن
# ═══════════════════════════════════════════════════════════
def _css(t):
    dark = t != "light"
    if dark:
        ac="#4ECBA0"; bg="linear-gradient(135deg,#080816 0%,#0d1526 50%,#080f1a 100%)"
        card="rgba(78,203,160,.08)"; brd="rgba(78,203,160,.25)"
        txt="#F0F4FF"; sub="rgba(78,203,160,.8)"; sbg="rgba(8,8,22,.97)"
        inp_bg="#111827"; inp_txt="#F0F4FF"
        code_bg="rgba(0,0,0,.45)"; code_txt="#7DFFC3"
    else:
        ac="#1a9e70"; bg="#EFF3F8"
        card="rgba(26,158,112,.07)"; brd="rgba(26,158,112,.25)"
        txt="#111827"; sub="rgba(26,158,112,.85)"; sbg="rgba(255,255,255,.99)"
        inp_bg="#FFFFFF"; inp_txt="#111827"
        code_bg="#F3F4F6"; code_txt="#065F46"
    tl="#111827" if not dark else "#94A3B8"
    return f"""
@import url('https://fonts.googleapis.com/css2?family=Space+Grotesk:wght@400;500;600;700&family=Tajawal:wght@400;500;700&display=swap');

/* ── Base ── */
.stApp {{
    background:{bg} !important;
    font-family:'Space Grotesk','Tajawal',system-ui,sans-serif !important;
}}
* {{ box-sizing:border-box; }}

/* ── Header ── */
.hdr {{ text-align:center; padding:1rem 0 .5rem; position:relative; }}
.hdr .brand-row {{
    display:flex; align-items:center; justify-content:center;
    gap:.5rem; margin-bottom:.3rem;
}}
.hdr .brand {{
    font-size:9px; font-weight:700; letter-spacing:.35em; color:{ac};
    text-transform:uppercase; opacity:.75;
}}
.hdr .dot {{ color:{ac}; font-size:7px; opacity:.5; }}
.hdr h1 {{
    font-family:'Space Grotesk',sans-serif; font-size:36px; font-weight:700;
    color:{txt}; margin:.1rem 0; letter-spacing:-.02em; line-height:1.1;
}}
.hdr h1 .ac {{ color:{ac}; }}
.hdr .glow {{
    display:inline-block;
    text-shadow: 0 0 40px {ac}55, 0 0 80px {ac}22;
}}
.hdr .tag-row {{
    display:flex; justify-content:center; gap:.5rem;
    margin:.4rem 0 .2rem; flex-wrap:wrap;
}}
.hdr .tag {{
    font-size:10px; font-weight:600;
    background:{card}; border:1px solid {brd};
    border-radius:20px; padding:3px 10px;
    color:{sub}; letter-spacing:.03em;
}}
.hdr .dv {{
    width:100px; height:2px;
    background:linear-gradient(90deg,transparent,{ac},{ac}88,transparent);
    margin:.4rem auto 0; border-radius:2px;
}}

/* ── Buttons ── */
.stButton>button {{
    background:linear-gradient(135deg,{ac} 0%,#1a9e6e 100%) !important;
    color:#051510 !important; font-weight:700 !important; font-size:14px !important;
    border:none !important; border-radius:10px !important;
    padding:.55rem 1.2rem !important;
    box-shadow:0 4px 15px {ac}40 !important;
    transition:all .2s ease !important;
    letter-spacing:.02em !important;
}}
.stButton>button:hover {{
    filter:brightness(1.1) !important;
    transform:translateY(-2px) !important;
    box-shadow:0 8px 25px {ac}55 !important;
}}
.stButton>button:active {{ transform:translateY(0) !important; }}

/* ── Text Areas — خط واضح وكبير ── */
textarea, .stTextArea textarea {{
    background:{inp_bg} !important;
    color:{inp_txt} !important;
    font-size:17px !important;
    font-family:'Tajawal','Space Grotesk',Arial,sans-serif !important;
    font-weight:600 !important;
    line-height:1.8 !important;
    border:2px solid {brd} !important;
    border-radius:12px !important;
    padding:14px 16px !important;
    caret-color:{ac} !important;
    letter-spacing:.01em !important;
    -webkit-font-smoothing:antialiased !important;
    text-shadow: none !important;
}}
textarea:focus, .stTextArea textarea:focus {{
    border-color:{ac} !important;
    box-shadow:0 0 0 3px {ac}28 !important;
    outline:none !important;
    color:{inp_txt} !important;
}}
textarea::placeholder, .stTextArea textarea::placeholder {{
    color:{sub} !important;
    opacity:.45 !important;
    font-weight:400 !important;
    font-size:15px !important;
}}
/* label الـ textarea */
.stTextArea label {{
    color:{sub} !important;
    font-size:11px !important;
    font-weight:700 !important;
    text-transform:uppercase !important;
    letter-spacing:.1em !important;
}}

/* ── Cards ── */
.card {{
    background:{card};
    border:1px solid {brd};
    border-radius:14px;
    padding:.9rem 1.1rem;
    margin:.5rem 0;
    position:relative;
    overflow:hidden;
}}
.card::before {{
    content:'';
    position:absolute; top:0; left:0; right:0; height:2px;
    background:linear-gradient(90deg,transparent,{ac},transparent);
}}
.card .lbl {{
    font-size:9px; font-weight:700; text-transform:uppercase;
    color:{sub}; letter-spacing:.2em;
}}
.card .body {{
    font-size:16px; color:{txt}; margin-top:.4rem; line-height:1.75;
    font-family:'Tajawal','Space Grotesk',sans-serif; font-weight:500;
}}
.card .meta {{
    font-size:10px; color:{sub}; margin-top:6px; opacity:.75;
    display:flex; gap:.6rem; flex-wrap:wrap;
}}

/* ── Code blocks — خط واضح ── */
.stCode, pre, code {{
    background:{code_bg} !important;
    color:{code_txt} !important;
    font-size:14px !important;
    font-family:'Space Grotesk',monospace !important;
    border:1px solid {brd} !important;
    border-radius:10px !important;
    padding:12px 16px !important;
    line-height:1.7 !important;
}}

/* ── Section headings ── */
.sh {{
    font-size:9px; font-weight:700; text-transform:uppercase;
    color:{sub}; margin:.65rem 0 .3rem; letter-spacing:.12em;
    display:flex; align-items:center; gap:.4rem;
}}
.sh::before {{
    content:''; display:inline-block;
    width:3px; height:12px;
    background:{ac}; border-radius:2px;
}}

/* ── Sidebar ── */
[data-testid="stSidebar"] {{
    background:{sbg} !important;
    border-right:1px solid {brd} !important;
}}

/* ── History items ── */
.hist {{
    padding:6px 10px; border-bottom:1px solid {brd}33;
    transition:background .15s;
}}
.hist:hover {{ background:{card}; border-radius:6px; }}
.hist .o {{ color:{txt}; font-size:12px; line-height:1.4; }}
.hist .tr {{ color:{ac}; font-size:12px; line-height:1.4; }}
.hist .m {{ font-size:9px; color:{sub}; opacity:.6; }}

/* ── Audio Input ── */
div[data-testid="stAudioInput"]>div {{
    background:{card} !important;
    border:2px solid {brd} !important;
    border-radius:60px !important;
    transition:border-color .2s !important;
}}
div[data-testid="stAudioInput"]>div:hover {{
    border-color:{ac} !important;
    box-shadow:0 0 20px {ac}30 !important;
}}

/* ── Group Chat Bubbles ── */
.cb {{
    border-radius:16px; padding:.9rem 1.2rem; margin-bottom:.8rem;
    border:1px solid {brd}; background:{card};
    transition:all .2s; position:relative; overflow:hidden;
}}
.cb::before {{
    content:''; position:absolute; top:0; left:0; bottom:0; width:3px;
}}
.cb:hover {{ transform:translateX(3px); box-shadow:0 6px 24px {brd}; }}
.cb-hd {{ display:flex; align-items:center; gap:10px; margin-bottom:.5rem; }}
.cb-ic {{ font-size:26px; line-height:1; filter:drop-shadow(0 2px 4px {brd}); }}
.cb-nfo {{ display:flex; flex-direction:column; gap:2px; }}
.cb-spk {{ font-size:14px; font-weight:700; }}
.cb-lang {{ font-size:10px; color:{sub}; opacity:.85; }}
.cb-ts {{ font-size:9px; color:{sub}; opacity:.5; }}
.cb-orig {{
    font-size:13px; color:{sub}; font-style:italic;
    border-left:3px solid {brd}; padding-left:10px;
    margin-bottom:.4rem; opacity:.9; line-height:1.5;
}}
.cb-trans {{
    font-size:16px; color:{txt}; font-weight:600; line-height:1.7;
    font-family:'Tajawal','Space Grotesk',sans-serif;
}}
.cb-eng {{ font-size:9px; color:{sub}; opacity:.5; margin-top:5px; }}

/* ── Tabs ── */
button[data-baseweb="tab"] {{
    font-family:'Space Grotesk',sans-serif !important; font-size:12px !important;
    font-weight:600 !important; color:{tl} !important;
    background:transparent !important; border:none !important;
    border-radius:8px 8px 0 0 !important; padding:.5rem 1rem !important;
    transition:all .2s !important;
}}
button[data-baseweb="tab"][aria-selected="true"] {{
    background:{card} !important; color:{ac} !important;
    border-bottom:2px solid {ac} !important;
}}
div[data-baseweb="tab-list"] {{
    gap:4px !important; border-bottom:1px solid {brd} !important;
}}

/* ── Selectbox ── */
.stSelectbox label, .stMultiSelect label {{ color:{sub} !important; font-size:11px !important; }}
.stSelectbox>div>div, .stMultiSelect>div>div {{
    background:{inp_bg} !important; color:{inp_txt} !important;
    border-color:{brd} !important; border-radius:8px !important;
    font-size:14px !important;
}}

/* ── Dividers ── */
hr {{
    margin:.5rem 0; border:none; height:1px;
    background:linear-gradient(90deg,transparent,{brd},transparent);
}}

/* ── Upload area ── */
[data-testid="stFileUploadDropzone"] {{
    background:{card} !important; border:1.5px dashed {brd} !important;
    border-radius:10px !important;
}}
[data-testid="stFileUploadDropzone"]:hover {{
    border-color:{ac} !important;
}}

/* ── Progress bar ── */
.stProgress > div > div {{ background:linear-gradient(90deg,{ac},{ac}99) !important; border-radius:4px !important; }}
"""

st.markdown(f"<style>{_css(st.session_state.theme)}</style>", unsafe_allow_html=True)

# ═══════════════════════════════════════════════════════════
#  Header
# ═══════════════════════════════════════════════════════════
st.markdown("""
<div class="hdr">
  <div class="brand-row">
    <span class="dot">◆</span>
    <span class="brand">Smart Voice Translator</span>
    <span class="dot">◆</span>
  </div>
  <h1 class="glow">HN <span class="ac">TRANSLATOR</span></h1>
  <div class="tag-row">
    <span class="tag">🎤 Voice</span>
    <span class="tag">🌐 Translate</span>
    <span class="tag">📷 Vision</span>
    <span class="tag">👥 Group</span>
  </div>
  <div class="dv"></div>
</div>
""", unsafe_allow_html=True)

# ═══════════════════════════════════════════════════════════
#  Sidebar
# ═══════════════════════════════════════════════════════════
with st.sidebar:
    if st.button("🌓  تبديل المظهر", use_container_width=True):
        st.session_state.theme = "light" if st.session_state.theme=="dark" else "dark"
        st.rerun()
    st.divider()

    with st.expander("🔑  مفاتيح API", expanded=False):
        for key,label,hint in [
            ("groq_api_key",   "🎤  Groq",  "STT + رؤية الصور — console.groq.com"),
            ("deepl_api_key",  "🌐  DeepL", "ترجمة عالية الجودة — 500K حرف/شهر"),
            ("cohere_api_key", "🔄  Cohere","STT احتياطي"),
        ]:
            nv=st.text_input(label,type="password",
                             value=st.session_state[key],help=hint,key=f"_i_{key}")
            if nv!=st.session_state[key]: st.session_state[key]=nv
        st.divider()
        if st.session_state.groq_api_key:
            st.success("🎙️ Groq · STT + Vision · نشط")
        else:
            st.warning("🎙️ Groq غير نشط — console.groq.com")
        if st.session_state.deepl_api_key:
            st.success("🌐 DeepL · ترجمة عالية الجودة · نشط")
        else:
            st.info("🌐 DeepL غير نشط · يُستخدم Google مجاناً")

    st.divider()
    history=get_history(80)
    if history:
        ca,cb=st.columns(2)
        with ca:
            if st.button("🗑️  مسح",use_container_width=True): clear_history(); st.rerun()
        with cb:
            b64h=base64.b64encode(json.dumps(history,ensure_ascii=False,indent=2).encode()).decode()
            st.markdown(f'<a href="data:application/json;base64,{b64h}" download="history.json">'
                        f'📥 تصدير</a>', unsafe_allow_html=True)
        for it in history:
            st.markdown(f"""<div class="hist">
<div class="o">{it.get('original','')[:55]}</div>
<div class="tr">{it.get('translated','')[:55]}</div>
<div class="m">{it.get('engine','')} · {it.get('target_lang','')} · {it.get('time','')}</div>
</div>""", unsafe_allow_html=True)
    else:
        st.markdown("<div style='text-align:center;font-size:30px;opacity:.2;padding:1.2rem'>📭</div>",
                    unsafe_allow_html=True)

# ═══════════════════════════════════════════════════════════
#  اختيار اللغات  ✅ زر التبديل مُصلح
# ═══════════════════════════════════════════════════════════
def _do_swap():
    """callback يعمل قبل إعادة الرسم — يضمن تبديل صحيح"""
    s = st.session_state.get("src_lang","Auto-Detect")
    t = st.session_state.get("tgt_lang","Arabic")
    new_src = t                                      # اللغة الهدف تصبح المصدر
    new_tgt = s if s != "Auto-Detect" else "Arabic"  # المصدر يصبح الهدف
    # تأكد أنهما مختلفتان
    if new_tgt == new_src:
        new_tgt = next((k for k in LANG_NAMES_NO_AUTO if k != new_src), "English")
    st.session_state["src_lang"] = new_src
    st.session_state["tgt_lang"] = new_tgt

st.markdown('<div class="sh">Translation Direction</div>', unsafe_allow_html=True)
c1, c2, c3 = st.columns([1, .2, 1])

with c1:
    # key="src_lang" يربط الـ selectbox بـ session_state مباشرةً
    src_name = st.selectbox("From", LANG_NAMES, key="src_lang")

with c2:
    st.markdown("<div style='height:24px'></div>", unsafe_allow_html=True)
    st.button("⇄", on_click=_do_swap, use_container_width=True, key="_swap_btn")

with c3:
    tgt_opts=[k for k in LANG_NAMES_NO_AUTO if k != src_name]
    # صحّح قيمة tgt_lang إن لم تعد في القائمة بعد التبديل
    if st.session_state.get("tgt_lang") not in tgt_opts:
        st.session_state["tgt_lang"] = tgt_opts[0]
    tgt_name = st.selectbox("To", tgt_opts, key="tgt_lang")

src_google = ALL_LANGUAGES.get(src_name,{}).get("google","auto")
tgt_gtts   = ALL_LANGUAGES.get(tgt_name,{}).get("gtts","en")

# ═══════════════════════════════════════════════════════════
#  التبويبات
# ═══════════════════════════════════════════════════════════
tab1,tab2,tab3,tab4,tab5 = st.tabs(
    ["🎙️ Voice","✍️ Text","📄 File","📸 Camera","🌍 Group Chat"])

# ─── Tab 1: Voice ───────────────────────────────────────────
with tab1:
    st.markdown("---")
    audio_val = st.audio_input("🎙️ اضغط وسجّل", key="_mic1", label_visibility="visible")
    if audio_val:
        with st.spinner("⏳  جاري التعرف على الكلام..."):
            recognized,_ = stt(audio_val.getvalue(), src_google)
        if recognized:
            st.success(f"✅  {recognized}")
            with st.spinner("⏳  الترجمة..."):
                translated,engine = translate_text(recognized, tgt_name, src_name)
            if translated:
                emotion = analyze_emotion(recognized)
                domains = detect_dom(recognized)
                d_str   = "  ".join(f'{_DOM[d]["e"]} {_DOM[d]["n"]}' for d in domains)
                st.markdown(f"""<div class="card">
<span class="lbl">✦  الترجمة</span>
<div class="body">{translated}</div>
<div class="meta">
  <span>{emotion}</span><span>·</span><span>{engine}</span>
  {'<span>·</span><span>'+d_str+'</span>' if d_str else ''}
</div></div>""", unsafe_allow_html=True)
                st.code(translated, language=None)
                aud=tts(translated,tgt_gtts)
                if aud: st.audio(aud,format="audio/mp3")
                save_translation(recognized,translated,emotion,src_name,tgt_name,engine)
            else:
                st.error(f"❌  {engine}")
        else:
            st.error("❌  تعذر التعرف على الكلام")

# ─── Tab 2: Text ────────────────────────────────────────────
with tab2:
    st.markdown("---")
    input_text = st.text_area(
        "📝  النص",
        height=120,
        placeholder="اكتب أو الصق النص هنا …",
        value=st.session_state.input_text,
        key="_txt"
    )
    st.session_state.input_text = input_text

    if input_text.strip():
        domains=detect_dom(input_text)
        if domains:
            b="".join(f'<span style="background:rgba(78,203,160,.13);border:1px solid '
                      f'rgba(78,203,160,.3);border-radius:20px;padding:2px 10px;'
                      f'font-size:11px;margin-right:5px">{_DOM[d]["e"]} {_DOM[d]["n"]}</span>'
                      for d in domains)
            st.markdown(f'<div style="margin-bottom:.5rem">{b}</div>',unsafe_allow_html=True)

    if st.button("Translate  ✦", use_container_width=True, key="_txt_btn"):
        if not input_text.strip():
            st.warning("الرجاء إدخال نص.")
        else:
            with st.spinner("جاري الترجمة..."):
                translated,engine=translate_text(input_text,tgt_name,src_name)
            if translated:
                emotion=analyze_emotion(input_text)
                st.markdown(f"""<div class="card">
<span class="lbl">✦  الترجمة</span>
<div class="body">{translated}</div>
<div class="meta"><span>{emotion}</span><span>·</span><span>{engine}</span></div>
</div>""", unsafe_allow_html=True)
                st.code(translated,language=None)
                aud=tts(translated,tgt_gtts)
                if aud: st.audio(aud,format="audio/mp3")
                save_translation(input_text,translated,emotion,src_name,tgt_name,engine)
            else:
                st.error(f"❌  {engine}")

# ─── Tab 3: File ────────────────────────────────────────────
with tab3:
    st.markdown("---")
    uploaded=st.file_uploader("اختر ملفاً  (PDF · DOCX · XLSX · TXT)",key="_file")
    if uploaded:
        st.caption(f"📎  {uploaded.name}  —  {len(uploaded.getvalue())//1024} KB")
        if st.button("🔍  استخراج وترجمة",key="_file_btn"):
            with st.spinner("استخراج النص..."):
                extracted,err=extract_file(uploaded.getvalue(),uploaded.name)
            if extracted:
                st.markdown('<div class="sh">النص المستخرج</div>',unsafe_allow_html=True)
                st.code(extracted[:2500]+("…" if len(extracted)>2500 else ""),language=None)
                st.caption(f"الكلمات: {len(extracted.split())}")
                with st.spinner("الترجمة..."):
                    translated,engine=translate_text(extracted,tgt_name,src_name)
                if translated:
                    emotion=analyze_emotion(extracted[:500])
                    st.markdown(f"""<div class="card">
<span class="lbl">✦  الترجمة</span>
<div class="body">{translated}</div>
<div class="meta"><span>{engine}</span></div>
</div>""", unsafe_allow_html=True)
                    ca,cb=st.columns(2)
                    with ca: st.download_button("📥  النص الأصلي",data=extracted,
                                                file_name="original.txt",mime="text/plain")
                    with cb: st.download_button("📥  الترجمة",data=translated,
                                                file_name="translated.txt",mime="text/plain")
                    aud=tts(translated,tgt_gtts)
                    if aud: st.audio(aud,format="audio/mp3")
                    save_translation(extracted[:300],translated,emotion,"File",tgt_name,engine)
                else:
                    st.error(f"❌  {engine}")
            else:
                st.error(f"❌  {err}")

# ─── Tab 4: Camera ──────────────────────────────────────────
with tab4:
    st.markdown("---")
    cam_file=st.file_uploader("📸 اختر صورة أو التقطها",type=["png","jpg","jpeg","webp","bmp"],key="_cam")
    if cam_file:
        img_bytes=cam_file.getvalue()
        try: st.image(Image.open(io.BytesIO(img_bytes)).convert("RGB"),
                      caption=cam_file.name,use_container_width=True)
        except Exception: pass

        if st.button("🔍  استخراج النص وترجمته",key="_cam_btn"):
            with st.spinner("جاري استخراج النص..."):
                ocr_langs=ALL_LANGUAGES.get(src_name,{}).get("ocr",["en"])
                extracted,err=ocr_image(img_bytes,ocr_langs)
            if not extracted:
                st.error(f"❌  {err}")
            else:
                lc,ln=detect_lang(extracted)
                src_for=_LC2NAME.get(lc,src_name)
                st.markdown('<div class="sh">النص المستخرج</div>',unsafe_allow_html=True)
                st.code(extracted,language=None)
                st.caption(f"اللغة المكتشفة: {ln}  ·  الكلمات: {len(extracted.split())}")
                with st.spinner("الترجمة..."):
                    translated,engine=translate_text(extracted,tgt_name,src_for)
                if translated:
                    emotion=analyze_emotion(extracted[:500])
                    st.markdown(f"""<div class="card">
<span class="lbl">✦  الترجمة  ({src_for} → {tgt_name})</span>
<div class="body">{translated}</div>
<div class="meta"><span>{emotion}</span><span>·</span><span>{engine}</span></div>
</div>""", unsafe_allow_html=True)
                    st.code(translated,language=None)
                    ca,cb=st.columns(2)
                    with ca: st.download_button("📥  النص",data=extracted,
                                                file_name="ocr.txt",mime="text/plain")
                    with cb: st.download_button("📥  الترجمة",data=translated,
                                                file_name="ocr_tr.txt",mime="text/plain")
                    aud=tts(translated,tgt_gtts)
                    if aud: st.audio(aud,format="audio/mp3")
                    save_translation(extracted,translated,emotion,f"Camera/{ln}",tgt_name,engine)
                else:
                    st.error(f"❌  {engine}")

# ─── Tab 5: Group Chat ──────────────────────────────────────
with tab5:
    st.markdown("---")

    st.markdown('<div class="sh">🎯 اللغة الهدف للترجمة</div>',unsafe_allow_html=True)
    grp_tgt=st.selectbox("اللغة الهدف",options=LANG_NAMES_NO_AUTO,
                         index=LANG_NAMES_NO_AUTO.index(tgt_name)
                         if tgt_name in LANG_NAMES_NO_AUTO else 0,
                         key="_g_tgt")
    grp_gtts=ALL_LANGUAGES.get(grp_tgt,{}).get("gtts","en")

    st.markdown("---")
    grp_audio=st.audio_input("🎤  سجّل المحادثة",key="_g_mic",label_visibility="visible")

    if grp_audio:
        if st.button("🚀  تحليل وترجمة المحادثة",use_container_width=True,key="_g_btn"):

            bar=st.progress(0,text="⏳  جاري التعرف وتمييز المتحدثين...")
            results,err=group_analyze(grp_audio.getvalue(),grp_tgt)
            bar.progress(80,text="⏳  الترجمة...")

            if not results:
                bar.empty(); st.error(f"❌  {err or 'تعذر تحليل الصوت'}")
            else:
                bar.progress(100,text="✅  اكتمل!"); time.sleep(0.3); bar.empty()
                spks=sorted({r["speaker_num"] for r in results})
                st.markdown(f"""<div style='text-align:center;margin:.5rem 0 1rem;
font-size:13px;opacity:.65'>
{len(results)} مقطع كلامي · {len(spks)} متحدث{"ون" if len(spks)>1 else ""}
</div>""", unsafe_allow_html=True)

                full_log=[]
                for item in results:
                    spk   = item["speaker_num"]
                    icon  = item["icon"]
                    lang  = item["lang"]
                    orig  = item["orig"]
                    trans = item["trans"] or "—"
                    eng   = item["engine"]
                    t0,t1 = item["t0"],item["t1"]
                    ts    = f"{t0:.1f}s – {t1:.1f}s" if t1>0 else ""

                    # لون فريد لكل متحدث
                    hue=(spk*83+140)%360
                    spk_col=f"hsl({hue},65%,62%)"
                    border_col=f"hsl({hue},65%,40%)"

                    st.markdown(f"""<div class="cb" style="border-left:3px solid {border_col}">
<div class="cb-hd">
  <span class="cb-ic">{icon}</span>
  <div class="cb-nfo">
    <span class="cb-spk" style="color:{spk_col}">المتحدث {spk}</span>
    <span class="cb-lang">🌐 {lang}</span>
    {'<span class="cb-ts">⏱ '+ts+'</span>' if ts else ''}
  </div>
</div>
<div class="cb-orig">"{orig}"</div>
<div class="cb-trans">{trans}</div>
<div class="cb-eng">{eng}</div>
</div>""", unsafe_allow_html=True)

                    aud=tts(trans,grp_gtts)
                    if aud: st.audio(aud,format="audio/mp3")

                    emotion=analyze_emotion(orig)
                    save_translation(orig,trans,emotion,f"Group/{lang}",grp_tgt,eng)
                    full_log.append(f"المتحدث {spk} ({lang}) [{ts}]\n"
                                    f"  الأصلي  : {orig}\n"
                                    f"  الترجمة : {trans}\n"
                                    f"  المشاعر : {emotion}\n")

                st.markdown("---")
                st.download_button("📥  تحميل محضر المحادثة",
                                   data="\n".join(full_log),
                                   file_name="group_transcript.txt",
                                   mime="text/plain",
                                   use_container_width=True,
                                   key="_g_dl")

# ═══════════════════════════════════════════════════════════
#  Footer
# ═══════════════════════════════════════════════════════════
st.markdown("""<div style='text-align:center;padding:1.5rem 0 .4rem;
color:rgba(100,130,170,.25);font-size:9px;letter-spacing:.14em;text-transform:uppercase'>
HN TRANSLATOR · Groq · DeepL · Google · Multi-Language Voice Suite
</div>""", unsafe_allow_html=True)
